"""
Attendance follow-ups and tools (15 Oct 2026):

  shifts/schedule/            yearly shift rota (Shift Schedule) + Excel
  shifts/on-duty/             who is rostered on each shift on a day
  shifts/cd-due/              people who worked on a gazetted holiday -> give 1 CD
  shifts/unplanned-leave/     leave marked on the sheet with no application -> raise one
  shifts/todo/                what is waiting for the signed-in person
  overtime/calculator/        try the overtime rules on a case (and the SOP worked examples)
  shifts/admin/sections/<id>/spread-groups/   place groups evenly over the rotation
"""
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
import calendar
import io

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from apps.core.utils import success, error, log_action
from . import access, roster, sheets as S
from .codes import CODES, kind
from .models import Section, DailySheet, AttendanceEntry, ExtraHours, ModuleSettings
from .overtime import day_rates, callout_rates

SHIFT_NAMES = {'M': 'Morning', 'E': 'Evening', 'N': 'Night', 'G': 'General', 'O': 'Off'}


def _section_from(request):
    sid = request.query_params.get('section') or request.data.get('section')
    if not sid:
        return None
    try:
        return Section.objects.select_related('department').get(pk=sid)
    except Exception:
        return None


def _month(request):
    raw = request.query_params.get('month') or request.data.get('month') or timezone.localdate().strftime('%Y-%m')
    try:
        y, m = [int(x) for x in str(raw).split('-')[:2]]
        return date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])
    except Exception:
        return None


def _schedule_sections(user):
    """Sections whose rota the person may see: attendance roles, plus Supervisors,
    SICs and Rota Managers from Leave (the rota is not personal data)."""
    if access.is_admin(user) or getattr(user, 'role', '') in ('manager', 'rota_manager'):
        return Section.objects.filter(is_active=True).select_related('department')
    secs = access.sections_for(user, 'attendance', 'view')
    qs = Section.objects.filter(is_active=True).select_related('department')
    if secs == access.ALL:
        return qs
    me = access.employee_of(user)
    mine = roster.membership_of(me, timezone.localdate()) if me else None
    ids = set(secs) | ({mine.group.section_id} if mine else set())
    return qs.filter(pk__in=ids)


class ScheduleSectionsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success([{'id': str(s.pk), 'name': str(s), 'pattern': s.pattern_list}
                        for s in _schedule_sections(request.user)])


class ScheduleView(APIView):
    """GET ?section=&year=[&export=xlsx]"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        section = _section_from(request)
        if section is None:
            return error('Choose a section.')
        if not _schedule_sections(request.user).filter(pk=section.pk).exists():
            return error('You do not have access to this section.', status=403)
        try:
            year = int(request.query_params.get('year') or timezone.localdate().year)
        except ValueError:
            return error('Year must be a number.')
        data = roster.year_schedule(section, year)
        data['section'] = str(section)
        data['gh'] = sorted(str(d) for d in S.gh_dates(date(year, 1, 1), date(year, 12, 31)))
        if request.query_params.get('export') == 'xlsx':
            return self._xlsx(section, year, data)
        return success(data)

    def _xlsx(self, section, year, data):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment
        fills = {'M': 'FFF4C2', 'E': 'D6E9FF', 'N': 'D9D2F3', 'G': 'DDF3E4', 'O': 'EEEEEE'}
        gh = set(data['gh'])
        wb = Workbook()
        wb.remove(wb.active)
        for m in range(1, 13):
            ws = wb.create_sheet(calendar.month_abbr[m])
            days = calendar.monthrange(year, m)[1]
            ws.cell(1, 1, f'{section} - shift schedule {calendar.month_name[m]} {year}').font = Font(bold=True)
            ws.cell(2, 1, 'Group').font = Font(bold=True)
            for d in range(1, days + 1):
                c = ws.cell(2, d + 1, f'{d}\n{calendar.day_abbr[date(year, m, d).weekday()][:2]}')
                c.alignment = Alignment(horizontal='center', wrap_text=True)
                c.font = Font(bold=True, color='C00000' if str(date(year, m, d)) in gh else '000000')
                ws.column_dimensions[c.column_letter].width = 4.5
            offset = (date(year, m, 1) - date(year, 1, 1)).days
            for i, row in enumerate(data['rows']):
                ws.cell(3 + i, 1, row['group']).font = Font(bold=True)
                for d in range(days):
                    code = row['shifts'][offset + d]
                    c = ws.cell(3 + i, d + 2, code)
                    c.alignment = Alignment(horizontal='center')
                    c.fill = PatternFill('solid', fgColor=fills.get(code, 'FFFFFF'))
            r = 4 + len(data['rows'])
            ws.cell(r, 1, 'M Morning  E Evening  N Night  G General  O Off   (red date = gazetted holiday)')
            ws.row_dimensions[2].height = 30
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="shift-schedule-{year}.xlsx"'
        return resp


class OnDutyView(APIView):
    """GET ?section=&date= : groups and people rostered on each shift."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        section = _section_from(request)
        if section is None:
            return error('Choose a section.')
        if not _schedule_sections(request.user).filter(pk=section.pk).exists():
            return error('You do not have access to this section.', status=403)
        try:
            day = date.fromisoformat(request.query_params.get('date') or str(timezone.localdate()))
        except ValueError:
            return error('Choose a valid date.')
        duty = roster.on_duty(section, day)
        types = S.shift_types(section)
        order = S.working_shifts(section) + ['O']
        shifts = []
        for code in order:
            t = types.get(code)
            groups = duty.get(code, {})
            shifts.append({'code': code, 'name': t.name if t else SHIFT_NAMES.get(code, code),
                           'start': str(t.start_time)[:5] if t else '', 'end': str(t.end_time)[:5] if t else '',
                           'groups': [{'group': g, 'people': ppl} for g, ppl in sorted(groups.items())]})
        return success({'section': str(section), 'date': str(day),
                        'is_gh': day in S.gh_dates(day, day), 'shifts': shifts})


# ------------------------------------------------------------------ CD for gazetted holidays

def _gh_holidays(start, end):
    from apps.leaves.models import PublicHoliday, HolidayCalendar
    qs = PublicHoliday.objects.filter(date__range=(start, end), is_optional=False)
    if HolidayCalendar.objects.filter(is_default=True).exists():
        qs = qs.filter(calendar__is_default=True)
    return {h.date: h for h in qs}


def cd_due(section, start, end):
    """People who worked on a gazetted holiday (locked days) and have no CD for it yet."""
    from apps.leaves.models import HolidayCompensation
    hols = _gh_holidays(start, end)
    if not hols:
        return []
    worked = {}
    for e in (AttendanceEntry.objects.filter(sheet__section=section, sheet__status='locked', sheet__date__in=list(hols))
              .select_related('employee', 'sheet')):
        if kind(e.code) == 'present' and e.shift_code != 'O':
            worked[(e.employee_id, e.sheet.date)] = (e.employee, f'{SHIFT_NAMES.get(e.shift_code, e.shift_code)} shift')
    for x in (ExtraHours.objects.filter(sheet__section=section, sheet__status='locked', sheet__date__in=list(hols))
              .select_related('employee', 'sheet')):
        worked.setdefault((x.employee_id, x.sheet.date), (x.employee, f'{x.hours:g} extra hours'))
    given = set(HolidayCompensation.objects.filter(holiday__in=list(hols.values()))
                .values_list('employee_id', 'holiday__date'))
    out = []
    for (eid, d), (emp, what) in sorted(worked.items(), key=lambda kv: (kv[0][1], kv[1][0].full_name)):
        if (eid, d) in given:
            continue
        h = hols[d]
        out.append({'employee_id': str(eid), 'name': emp.full_name, 'p_number': emp.p_number or '',
                    'date': str(d), 'holiday_id': str(h.pk), 'holiday': h.name, 'worked': what})
    return out


class CdDueView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        section = _section_from(request)
        rng = _month(request)
        if section is None or rng is None:
            return error('Choose a section and month.')
        if not access.can(request.user, 'attendance', 'review', section.pk):
            return error('Only the section Supervisor or Manager sees this list.', status=403)
        return success(cd_due(section, *rng))

    def post(self, request):
        from apps.employees.models import Employee
        from apps.leaves.models import PublicHoliday
        from apps.leaves.holidays import give_holiday_cd
        section = _section_from(request)
        if section is None:
            return error('Choose a section.')
        if not access.can(request.user, 'attendance', 'review', section.pk):
            return error('Only the section Supervisor or Manager can give this.', status=403)
        h = PublicHoliday.objects.filter(pk=request.data.get('holiday_id')).first() if request.data.get('holiday_id') else None
        emp = Employee.objects.filter(pk=request.data.get('employee_id')).first() if request.data.get('employee_id') else None
        if h is None or emp is None:
            return error('Choose the person and the holiday.')
        due = {(r['employee_id'], r['holiday_id']) for r in cd_due(section, h.date, h.date)}
        if (str(emp.pk), str(h.pk)) not in due:
            return error(f'{emp.full_name} is not shown as working on {h.name} in a locked sheet of this section.')
        me = access.employee_of(request.user)
        if me is not None and me.pk == emp.pk and not access.is_admin(request.user):
            return error('You cannot give yourself a compensatory day.', status=403)
        try:
            comp, bal = give_holiday_cd(emp, h, request.user, request.data.get('note') or 'From the attendance sheet')
        except ValueError as e:
            return error(str(e))
        log_action(request, 'attendance.holiday_cd', 'HolidayCompensation', comp.pk,
                   {'employee': emp.full_name, 'holiday': h.name, 'date': str(h.date), 'section': str(section)})
        return success({'id': str(comp.id), 'cd_available': str(bal.available)},
                       f'1 compensatory day added for {emp.full_name}.', status=201)


# ------------------------------------------------------------------ unplanned leave

PREFERRED_TYPE = {'V': 'AL', 'C': 'CL', 'S': 'SL', 'L': 'UL', 'BD': 'BD', 'CD': 'CD'}


def leave_type_for(code):
    from apps.leaves.models import LeaveType
    pref = PREFERRED_TYPE.get(code)
    lt = LeaveType.objects.filter(code=pref, is_active=True).first() if pref else None
    if lt:
        return lt
    cmap = S.leave_code_map()
    for t in LeaveType.objects.filter(is_active=True):
        if cmap.get((t.code or '').upper()) == code:
            return t
    return None


def unplanned(section, start, end):
    """Leave marked on the sheet with no leave application, grouped into runs of
    consecutive days per person and code."""
    rows = (AttendanceEntry.objects.filter(sheet__section=section, sheet__date__range=(start, end),
                                           leave_application__isnull=True)
            .select_related('employee', 'sheet').order_by('employee__full_name', 'sheet__date'))
    runs = []
    for e in rows:
        if kind(e.code) != 'leave':
            continue
        d = e.sheet.date
        last = runs[-1] if runs else None
        if last and last['employee_id'] == str(e.employee_id) and last['code'] == e.code \
                and date.fromisoformat(last['end']) + timedelta(days=1) == d:
            last['end'] = str(d)
            last['days'] += 1
            continue
        runs.append({'employee_id': str(e.employee_id), 'name': e.employee.full_name,
                     'p_number': e.employee.p_number or '', 'code': e.code,
                     'code_label': CODES.get(e.code, {}).get('label', e.code),
                     'start': str(d), 'end': str(d), 'days': 1, 'note': e.note})
    for r in runs:
        lt = leave_type_for(r['code'])
        r['leave_type'] = lt.name if lt else None
    return runs


class UnplannedLeaveView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        section = _section_from(request)
        rng = _month(request)
        if section is None or rng is None:
            return error('Choose a section and month.')
        if not access.can(request.user, 'attendance', 'review', section.pk):
            return error('Only the section Supervisor or Manager sees this list.', status=403)
        return success(unplanned(section, *rng))

    def post(self, request):
        """Raise a leave application for a run of unplanned leave. It then follows
        the normal leave approval (Supervisor, then SIC)."""
        from apps.employees.models import Employee
        from apps.leaves.services import LeaveApplicationService, LeaveValidationError
        section = _section_from(request)
        if section is None:
            return error('Choose a section.')
        if not access.can(request.user, 'attendance', 'review', section.pk):
            return error('Only the section Supervisor or Manager can do this.', status=403)
        emp = Employee.objects.filter(pk=request.data.get('employee_id')).first() if request.data.get('employee_id') else None
        code = str(request.data.get('code') or '')
        try:
            start = date.fromisoformat(str(request.data.get('start')))
            end = date.fromisoformat(str(request.data.get('end')))
        except ValueError:
            return error('Choose the dates.')
        if emp is None or end < start:
            return error('Choose the person and the dates.')
        match = [r for r in unplanned(section, start, end)
                 if r['employee_id'] == str(emp.pk) and r['code'] == code]
        if not match or match[0]['start'] != str(start) or match[0]['end'] != str(end):
            return error('These days are not shown as unplanned leave on the sheet any more. Refresh the list.')
        if code == 'CD':
            return error('A compensatory day needs the duty date it was earned for; please apply for it from Leave.')
        lt = leave_type_for(code)
        if lt is None:
            return error(f'No leave type is linked to code {code}. Admin can link it in Attendance settings.')
        me = access.employee_of(request.user)
        reason = str(request.data.get('reason') or '').strip() or \
            f'Marked as {CODES.get(code, {}).get("label", code)} on the attendance sheet' + \
            (f' by {me.full_name}' if me else '')
        try:
            app = LeaveApplicationService.submit_application(
                emp, lt, start, end, reason[:500], request=request,
                contact_during_leave='-', address_during_leave='-')
        except LeaveValidationError as e:
            return error(f'The leave could not be raised: {e}')
        AttendanceEntry.objects.filter(sheet__section=section, employee=emp, code=code,
                                       sheet__date__range=(start, end), leave_application__isnull=True
                                       ).update(leave_application=app)
        log_action(request, 'attendance.unplanned_leave_raised', 'LeaveApplication', app.pk,
                   {'employee': emp.full_name, 'code': code, 'start': str(start), 'end': str(end)})
        return success({'id': str(app.pk), 'reference': app.reference_number, 'status': app.status},
                       f'{lt.name} application {app.reference_number} raised for {emp.full_name}; '
                       'it now goes for approval.', status=201)


# ------------------------------------------------------------------ to-do

class TodoView(APIView):
    """What is waiting for the signed-in person in Attendance (last 31 days)."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from .models import ShiftSubmission, Correction
        today = timezone.localdate()
        start = today - timedelta(days=31)
        user = request.user

        def secs(action):
            s = access.sections_for(user, 'attendance', action)
            return None if s == access.ALL else set(s)

        def scope(qs, action, field='section_id'):
            s = secs(action)
            if s is None:
                return qs
            return qs.filter(**{f'{field}__in': s})

        out = {}
        mark = scope(ShiftSubmission.objects.filter(status='sent_back', sheet__date__gte=start), 'mark', 'sheet__section_id')
        out['sent_back'] = mark.count() if (secs('mark') is None or secs('mark')) else 0
        rev = scope(DailySheet.objects.filter(status='open', date__gte=start, date__lt=today), 'review')
        out['to_review'] = sum(1 for sh in rev.prefetch_related('submissions') if S.all_submitted(sh)) \
            if (secs('review') is None or secs('review')) else 0
        lock = scope(DailySheet.objects.filter(status='reviewed', date__gte=start, date__lt=today), 'lock')
        out['to_lock'] = lock.count() if (secs('lock') is None or secs('lock')) else 0
        dec = scope(Correction.objects.filter(status='pending'), 'decide', 'sheet__section_id')
        out['corrections'] = dec.count() if (secs('decide') is None or secs('decide')) else 0
        cd = unp = 0
        rsecs = secs('review')
        if rsecs is None or rsecs:
            qs = Section.objects.filter(is_active=True)
            if rsecs is not None:
                qs = qs.filter(pk__in=rsecs)
            for sec in qs:
                cd += len(cd_due(sec, start, today))
                unp += len(unplanned(sec, start, today))
        out['cd_due'] = cd
        out['unplanned_leave'] = unp
        out['total'] = sum(out.values())
        return success(out)


# ------------------------------------------------------------------ overtime calculator

WORKED_EXAMPLES = [
    {'title': '2 hours after the shift on a working day', 'day_type': 'WORK', 'kind': 'normal', 'hours': 2, 'crew_complete': False, 'expected': 3},
    {'title': '2 hours on a working day, 48-hour week already complete', 'day_type': 'WORK', 'kind': 'normal', 'hours': 2, 'crew_complete': True, 'expected': 4},
    {'title': '8 hours worked on a First Day Off', 'day_type': 'FO', 'kind': 'normal', 'hours': 8, 'crew_complete': False, 'expected': 16},
    {'title': '10 hours worked on a Second Day Off', 'day_type': 'SO', 'kind': 'normal', 'hours': 10, 'crew_complete': False, 'expected': 12},
    {'title': '8 hours overtime on a gazetted holiday', 'day_type': 'GH', 'kind': 'normal', 'hours': 8, 'crew_complete': False, 'expected': 24},
    {'title': 'Emergency call-out of 2 hours', 'day_type': 'WORK', 'kind': 'callout', 'hours': 2, 'crew_complete': False, 'expected': 4},
    {'title': 'Emergency call-out of 5 hours', 'day_type': 'WORK', 'kind': 'callout', 'hours': 5, 'crew_complete': False, 'expected': 9},
    {'title': 'Call-out of 3 hours on a gazetted holiday', 'day_type': 'GH', 'kind': 'callout', 'hours': 3, 'crew_complete': False, 'expected': 12},
]


def calculate(day_type, kind_, hours, crew_complete, shift, rules):
    h = Decimal(str(hours))
    if kind_ == 'callout':
        s, d, t = callout_rates(day_type == 'GH', h, rules)
        if day_type == 'GH':
            why = (f'Call-out on a gazetted holiday under {rules["callout_gh_short_below_hours"]:g} hours is paid as '
                   f'{rules["callout_gh_short_paid_hours"]:g} hours at single rate' if h < Decimal(str(rules['callout_gh_short_below_hours']))
                   else 'Call-out on a gazetted holiday: all hours at triple rate')
        elif h <= Decimal(str(rules['callout_short_max_hours'])):
            why = (f'Call-out of up to {rules["callout_short_max_hours"]:g} hours is paid as '
                   f'{rules["callout_short_paid_hours"]:g} hours at single rate')
        else:
            why = 'Call-out: first hour at single rate, the rest at double rate'
    else:
        s, d, t = day_rates(day_type, crew_complete, h, rules)
        if day_type == 'GH':
            why = 'Gazetted holiday: all overtime hours at triple rate'
        elif day_type == 'FO':
            why = 'First Day Off: all hours at double rate'
        elif day_type == 'SO':
            why = f'Second Day Off: first {rules["second_off_single_hours"]:g} hours at single rate, the rest at double rate'
        elif crew_complete:
            why = f'{rules["crew_week_hours"]:g}-hour crew week already complete: all hours at double rate'
        else:
            why = f'Working day: first {rules["normal_day_first_hours_single"]:g} hour(s) at single rate, the rest at double rate'
    eq = s + 2 * d + 3 * t
    evening = night = Decimal('0')
    if shift in ('E', 'N') and rules.get('premium_on_overtime', True):
        if shift == 'E':
            evening = h
        else:
            night = h
    return {'single': str(s), 'double': str(d), 'triple': str(t), 'equivalent_single': str(eq),
            'evening_premium_hours': str(evening), 'night_premium_hours': str(night),
            'evening_premium_equivalent': str((evening * Decimal(str(rules['evening_premium_percent'])) / 100).quantize(Decimal('0.01'))),
            'night_premium_equivalent': str((night * Decimal(str(rules['night_premium_percent'])) / 100).quantize(Decimal('0.01'))),
            'explanation': why}


def _can_use_calculator(user):
    return access.is_admin(user) or bool(access.sections_for(user, 'overtime', 'view')) \
        or bool(access.sections_for(user, 'attendance', 'view'))


class CalculatorView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not _can_use_calculator(request.user):
            return error('You do not have access to Overtime.', status=403)
        rules = ModuleSettings.get().rules()
        examples = []
        for ex in WORKED_EXAMPLES:
            r = calculate(ex['day_type'], ex['kind'], ex['hours'], ex['crew_complete'], '', rules)
            examples.append({**ex, 'result': r,
                             'matches_sop': Decimal(r['equivalent_single']) == Decimal(str(ex['expected']))})
        return success({'rules': rules, 'examples': examples})

    def post(self, request):
        if not _can_use_calculator(request.user):
            return error('You do not have access to Overtime.', status=403)
        dt = str(request.data.get('day_type') or 'WORK').upper()
        k = str(request.data.get('kind') or 'normal')
        if dt not in ('WORK', 'FO', 'SO', 'GH') or k not in ('normal', 'training', 'callout'):
            return error('Choose the day type and the kind of overtime.')
        try:
            hours = Decimal(str(request.data.get('hours')))
        except (InvalidOperation, TypeError):
            return error('Enter the hours.')
        if not hours.is_finite() or hours <= 0 or hours > 24:
            return error('Hours must be between 0 and 24.')
        crew = request.data.get('crew_complete') in (True, 'true', 'True', '1', 1)
        shift = str(request.data.get('shift') or '').upper()
        return success(calculate(dt, k, hours, crew, shift, ModuleSettings.get().rules()))


# ------------------------------------------------------------------ admin helper

class SpreadGroupsView(APIView):
    """POST shifts/admin/sections/<id>/spread-groups/ : place the groups evenly over
    the rotation (e.g. 4 groups on a 24-day rota start 0, 6, 12, 18 days apart)."""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        if not access.is_admin(request.user):
            return error('Permission denied.', status=403)
        section = Section.objects.filter(pk=pk).first()
        if section is None:
            return error('Section not found.', status=404)
        if DailySheet.objects.filter(section=section, status='locked').exists():
            return error('Days are already locked for this section, so the rotation cannot change. '
                         'Use roster changes for single days instead.')
        groups = list(section.groups.all().order_by('order', 'name'))
        if request.data.get('preset') == 'engro_2026':
            return self._engro_2026(request, section, groups)
        n = len(section.pattern_list or [])
        if not groups or not n:
            return error('Add the rotation and the groups first.')
        step = n // len(groups) or 1
        # Group A starts the cycle; each next group is `step` days ahead
        for i, g in enumerate(groups):
            g.offset = (i * step) % n
            g.save(update_fields=['offset'])
        log_action(request, 'shifts.spread_groups', 'Section', section.pk,
                   {'pattern': section.pattern, 'offsets': {g.name: g.offset for g in groups}})
        return success({'offsets': {g.name: g.offset for g in groups}}, 'Groups placed evenly over the rotation.')

    def _engro_2026(self, request, section, groups):
        """Engro's "Production Department Shift Schedule 2026" (shared by Engro on 16 Oct 2026):
        each crew works 6 Morning, 2 off, 6 Night, 2 off, 6 Evening, 2 off; cycle day 1 is
        1 Jan 2026 with A on Morning, B Evening, C Night, D Off."""
        names = {g.name.strip().upper(): g for g in groups}
        if set(names) != set(ENGRO_2026_OFFSETS):
            return error('This schedule needs exactly four groups named A, B, C and D.')
        section.pattern = ENGRO_2026_PATTERN
        section.pattern_anchor = date(2026, 1, 1)
        section.save(update_fields=['pattern', 'pattern_anchor'])
        for name, off in ENGRO_2026_OFFSETS.items():
            names[name].offset = off
            names[name].save(update_fields=['offset'])
        log_action(request, 'shifts.spread_groups', 'Section', section.pk,
                   {'preset': 'engro_2026', 'pattern': section.pattern, 'offsets': ENGRO_2026_OFFSETS})
        return success({'offsets': ENGRO_2026_OFFSETS}, "Engro's 2026 shift schedule applied to groups A-D.")


ENGRO_2026_PATTERN = ','.join(['M'] * 6 + ['O'] * 2 + ['N'] * 6 + ['O'] * 2 + ['E'] * 6 + ['O'] * 2)
# Position of each crew in the rotation on 1 Jan 2026 (A = 2nd Morning, D = 2nd off after
# Mornings, C = 6th Night, B = 4th Evening) - read from Engro's PDF / Excel.
ENGRO_2026_OFFSETS = {'A': 1, 'B': 19, 'C': 13, 'D': 7}
