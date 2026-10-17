"""Load Engro's religion list (Engro, 17 Oct 2026: religion-based gazetted
holidays go only to people of that religion; Engro sends the list).

The list is an Excel (.xlsx) or CSV file with a header row containing a
"Religion" column and one of "P.No" / "Employee ID" / "Name". Each row is matched
to an employee by P.No., then Employee ID (e.g. MPT-AMM2-07), then by name when
exactly one active person has that name. Only the religion field is changed;
rows that cannot be matched are reported, not guessed.

Assumptions (R-1..R-3, Engro to confirm):
  R-1 Words accepted: Muslim / Islam, Christian / Christianity, Hindu / Hinduism,
      Sikh / Sikhism; anything else non-empty is recorded as "Other".
  R-2 People not in the list keep what they have (blank = not recorded); a blank
      religion gets the holidays for everyone only.
  R-3 Eid, Ashura etc. are gazetted for everyone in Pakistan, so they stay "for
      everyone"; only holidays marked with a religion are limited.
"""
import csv
import io
import re

from django.db import transaction

from .organogram import name_key

WORDS = {'muslim': 'muslim', 'islam': 'muslim', 'christian': 'christian', 'christianity': 'christian',
         'hindu': 'hindu', 'hinduism': 'hindu', 'sikh': 'sikh', 'sikhism': 'sikh'}
LABEL = {'muslim': 'Muslim', 'christian': 'Christian', 'hindu': 'Hindu', 'sikh': 'Sikh', 'other': 'Other', '': 'not recorded'}


class ReligionImportError(Exception):
    pass


def religion_code(text):
    t = re.sub(r'[^a-z]', '', str(text or '').lower())
    if not t:
        return ''
    return WORDS.get(t, 'other')


def _table(fileobj, filename):
    if str(filename or '').lower().endswith('.csv'):
        raw = fileobj.read()
        text = raw.decode('utf-8-sig', errors='replace') if isinstance(raw, bytes) else raw
        return [list(r) for r in csv.reader(io.StringIO(text))]
    import openpyxl
    try:
        wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    except Exception:
        raise ReligionImportError('This is not an Excel (.xlsx) or CSV file.')
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if any(any('religion' in str(c or '').lower() for c in r) for r in rows[:15]):
            return rows
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def read_rows(fileobj, filename=''):
    rows = _table(fileobj, filename)
    head = cols = None
    for i, r in enumerate(rows[:15]):
        cells = [str(c or '').strip().lower() for c in r]
        if any('religion' in c for c in cells) and \
                any(k in c for c in cells for k in ('name', 'p.no', 'p no', 'pno', 'id')):
            head, cols = i, cells
            break
    if head is None:
        raise ReligionImportError('No header row with "Religion" and "P.No" / "Employee ID" / "Name" columns found.')

    def find(*keys):
        return next((j for j, c in enumerate(cols) if any(k in c for k in keys)), None)
    rel = find('religion')
    pno = find('p.no', 'p no', 'pno', 'p_no', 'personal')
    eid = find('employee id', 'emp id', 'employee_id', 'id')
    name = find('name')
    if eid == pno:
        eid = None
    if pno is None and eid is None and name is None:
        raise ReligionImportError('The file needs a "P.No", "Employee ID" or "Name" column next to "Religion".')

    def cell(r, j):
        if j is None or j >= len(r) or r[j] is None:
            return ''
        v = str(r[j]).strip()
        return str(int(float(v))) if re.fullmatch(r'\d+\.0', v) else v
    out = []
    for r in rows[head + 1:]:
        row = {'p_number': cell(r, pno), 'employee_id': cell(r, eid), 'name': ' '.join(cell(r, name).split()),
               'religion_text': cell(r, rel)}
        if not (row['p_number'] or row['employee_id'] or row['name']):
            continue
        row['religion'] = religion_code(row['religion_text'])
        out.append(row)
    return out


def plan(rows):
    from .models import Employee
    people = list(Employee.objects.all())
    by_pno = {e.p_number: e for e in people if e.p_number}
    by_id = {e.employee_id.lower(): e for e in people}
    by_name = {}
    for e in people:
        if e.status in ('active', 'on_leave'):
            by_name.setdefault(name_key(e.full_name), []).append(e)
    seen = {}
    result = []
    for row in rows:
        r = dict(row)
        emp, how = None, ''
        if row['p_number'] and row['p_number'] in by_pno:
            emp, how = by_pno[row['p_number']], 'P.No.'
        elif row['p_number'] and row['p_number'].lower() in by_id:
            emp, how = by_id[row['p_number'].lower()], 'Employee ID'
        elif row['employee_id'] and row['employee_id'].lower() in by_id:
            emp, how = by_id[row['employee_id'].lower()], 'Employee ID'
        elif row['name'] and not row['p_number'] and not row['employee_id']:
            same = by_name.get(name_key(row['name']), [])
            if len(same) > 1:
                r.update(action='skip', reason=f'{len(same)} employees have this name - add the P.No. or Employee ID')
                result.append(r)
                continue
            if same:
                emp, how = same[0], 'name'
        if emp is None:
            r.update(action='skip', reason='No employee with this P.No. / Employee ID / name')
        elif not row['religion']:
            r.update(action='skip', reason='Religion is empty', employee_id=str(emp.pk), current_name=emp.full_name)
        elif str(emp.pk) in seen:
            r.update(action='skip', reason=f'Same person as an earlier row ({seen[str(emp.pk)]})')
        else:
            seen[str(emp.pk)] = row['p_number'] or row['employee_id'] or row['name']
            r.update(employee_id=str(emp.pk), current_name=emp.full_name, matched_by=how,
                     current=emp.religion or '', current_label=LABEL.get(emp.religion or '', emp.religion),
                     new_label=LABEL[row['religion']],
                     action='same' if (emp.religion or '') == row['religion'] else 'update')
        result.append(r)
    return result


def summary(planned):
    s = {}
    for r in planned:
        s[r['action']] = s.get(r['action'], 0) + 1
    by = {}
    for r in planned:
        if r['action'] in ('update', 'same'):
            by[LABEL[r['religion']]] = by.get(LABEL[r['religion']], 0) + 1
    s['by_religion'] = by
    return s


@transaction.atomic
def apply(rows):
    from .models import Employee
    planned = plan(rows)
    for r in planned:
        if r['action'] == 'update':
            Employee.objects.filter(pk=r['employee_id']).update(religion=r['religion'])
    return planned
