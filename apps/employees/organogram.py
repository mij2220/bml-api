"""Load Engro's production organogram (Excel, sheet "Staff List") into the app.

Each row: Unit, Level, Name, Cadre / P.No, Role, List. For every person:
  * found (by P.No., otherwise by name) -> unit and place in the unit (Unit
    Manager / SIC / Supervisor-Engineer / Operator) are updated; approvers get
    the approver login role;
  * not found -> an employee record and a login are created.
Nothing is deleted, and nobody's Admin role is lowered.

Assumptions (Engro to confirm; all editable afterwards in Admin > Employees):
  O-1 Rows for "Department" (General Manager) and "PP" (Secretary) are not a
      production unit and are skipped.
  O-2 Managers, SICs and Supervisors have no P.No. in the sheet: they are matched
      by name; if new they get an ID like MPT-AMM2-01 until Engro sends P.Nos.
  O-3 New people: gender male, joining date 1 Jan 2026, permanent - to be
      corrected by Admin (marked "check details" in the result).
  O-4 Login: P.No. (or the MPT ID) with the password Employee@1234 / Manager@1234,
      changed at first sign-in. Email is a placeholder <id>@staff.shiftledger.local.
"""
import re
from datetime import date

from django.db import transaction

LEVEL_ROLE = {'unit manager': 'manager', 'sic': 'sic', 'engineer/supervisor': 'supervisor',
              'supervisor': 'supervisor', 'engineer': 'supervisor', 'operator': ''}
SKIP_UNITS = {'department', 'pp'}
APPROVER_ROLES = ('manager', 'sic', 'supervisor')
PLACEHOLDER_JOIN = date(2026, 1, 1)
EMAIL_DOMAIN = 'staff.shiftledger.local'


class OrganogramError(Exception):
    pass


def unit_key(name):
    """'Ammonia-2', 'AMM-2', 'Ammonia-II' -> 'amm2'; 'Urea-2' / 'UR-2' -> 'ur2'; 'PM&S' -> 'pm&s'."""
    k = re.sub(r'[\s\-_.]', '', str(name or '').lower())
    k = re.sub(r'iii$', '3', k)
    k = re.sub(r'ii$', '2', k)
    k = k.replace('ammonia', 'amm').replace('urea', 'ur')
    return k


def name_key(name):
    return re.sub(r'[^a-z ]', '', re.sub(r'[\s\-]+', ' ', str(name or '').lower())).strip()


def read_rows(fileobj):
    """Rows of the "Staff List" sheet as dicts."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    except Exception:
        raise OrganogramError('This is not an Excel file (.xlsx).')
    ws = wb['Staff List'] if 'Staff List' in wb.sheetnames else wb.worksheets[-1]
    rows = list(ws.iter_rows(values_only=True))
    head = None
    for i, r in enumerate(rows[:15]):
        cells = [str(c or '').strip().lower() for c in r]
        if 'unit' in cells and 'name' in cells:
            head = i
            cols = {c: j for j, c in enumerate(cells)}
            break
    if head is None:
        raise OrganogramError('No header row with "Unit" and "Name" found (expected sheet "Staff List").')
    pno_col = next((j for c, j in cols.items() if 'p.no' in c or 'cadre' in c), None)
    out = []
    for r in rows[head + 1:]:
        def g(key):
            j = cols.get(key)
            return r[j] if j is not None and j < len(r) else None
        name = str(g('name') or '').strip()
        if not name:
            continue
        raw = r[pno_col] if pno_col is not None and pno_col < len(r) else None
        raw = '' if raw is None else str(raw).strip()
        if re.fullmatch(r'\d+(\.0)?', raw):
            pno, cadre = str(int(float(raw))), ''
        else:
            pno, cadre = '', raw
        out.append({'unit': str(g('unit') or '').strip(), 'level': str(g('level') or '').strip(),
                    'name': ' '.join(name.split()), 'p_number': pno, 'cadre': cadre,
                    'role_text': str(g('role') or '').strip(), 'list': str(g('list') or '').strip()})
    return out


def plan(rows):
    """What loading the rows would do, without changing anything."""
    from apps.employees.models import Employee, Department
    units = {unit_key(d.name): d for d in Department.objects.all()}
    by_pno = {e.p_number: e for e in Employee.objects.exclude(p_number__isnull=True).exclude(p_number='')}
    by_name = {}
    for e in Employee.objects.all():
        by_name.setdefault(name_key(e.full_name), []).append(e)
    result = []
    for row in rows:
        r = dict(row)
        uk = unit_key(row['unit'])
        r['unit_role'] = LEVEL_ROLE.get(row['level'].lower(), '')
        if uk in SKIP_UNITS or not uk:
            r.update(action='skip', reason='Not a production unit')
            result.append(r)
            continue
        dept = units.get(uk)
        r['unit_exists'] = dept is not None
        r['unit_name'] = dept.name if dept else row['unit']
        emp, how = None, ''
        if row['p_number'] and row['p_number'] in by_pno:
            emp, how = by_pno[row['p_number']], 'P.No.'
        if emp is None:
            same = by_name.get(name_key(row['name']), [])
            if len(same) == 1:
                emp, how = same[0], 'name'
            elif len(same) > 1:
                r.update(action='skip', reason=f'{len(same)} employees have this name - fix by hand')
                result.append(r)
                continue
        if emp is not None:
            changes = []
            if dept is None or emp.department_id != dept.pk:
                changes.append(f"unit {emp.department.name if emp.department_id else '-'} -> {r['unit_name']}")
            if (emp.unit_role or '') != r['unit_role']:
                changes.append(f"place {emp.unit_role or 'staff'} -> {r['unit_role'] or 'staff'}")
            if r['unit_role'] in APPROVER_ROLES and emp.user.role == 'employee':
                changes.append('login role employee -> manager')
            r.update(action='update' if changes else 'same', employee_id=str(emp.pk),
                     current_name=emp.full_name, matched_by=how, changes=changes)
        else:
            r.update(action='create', reason='' if row['p_number'] else 'No P.No. in the sheet: a temporary MPT ID is given')
        result.append(r)
    return result


def summary(planned):
    s = {}
    for r in planned:
        s[r['action']] = s.get(r['action'], 0) + 1
    s['new_units'] = sorted({r['unit_name'] for r in planned if r.get('unit_exists') is False and r['action'] != 'skip'})
    return s


@transaction.atomic
def apply(rows, actor_user=None):
    """Load the rows. Returns the plan with what was done, plus logins created."""
    from apps.employees.models import Employee, Department, Designation
    from apps.accounts.models import User
    from apps.leaves.services import initialize_employee_balances
    from django.contrib.auth.hashers import make_password
    planned = plan(rows)
    units = {unit_key(d.name): d for d in Department.objects.all()}
    hashes = {}      # hash each first password once: hundreds of logins must not time out
    logins = []
    counters = {}
    for r in planned:
        if r['action'] in ('skip', 'same'):
            continue
        uk = unit_key(r['unit'])
        dept = units.get(uk)
        if dept is None:
            dept = Department.objects.create(name=r['unit'])
            units[uk] = dept
        role_text = r['role_text'] or {'manager': 'Unit Manager', 'sic': 'Shift Incharge',
                                       'supervisor': 'Shift Supervisor'}.get(r['unit_role'], 'Operator')
        desig, _ = Designation.objects.get_or_create(name=role_text[:100], department=dept)
        if r['action'] == 'update':
            emp = Employee.objects.select_related('user').get(pk=r['employee_id'])
            emp.department = dept
            emp.unit_role = r['unit_role']
            emp.save(update_fields=['department', 'unit_role', 'updated_at'])
            if r['unit_role'] in APPROVER_ROLES and emp.user.role == 'employee':
                emp.user.role = 'manager'
                emp.user.save(update_fields=['role'])
            continue
        # create
        if r['p_number']:
            emp_id = r['p_number']
        else:
            n = counters.get(uk, 0)
            while True:
                n += 1
                emp_id = f'MPT-{uk.upper().replace("&", "")}-{n:02d}'
                if not Employee.objects.filter(employee_id=emp_id).exists():
                    break
            counters[uk] = n
        if Employee.objects.filter(employee_id=emp_id).exists():
            r.update(action='skip', reason=f'Employee ID {emp_id} already used')
            continue
        approver = r['unit_role'] in APPROVER_ROLES
        password = 'Manager@1234' if approver else 'Employee@1234'
        email = f'{emp_id.lower()}@{EMAIL_DOMAIN}'
        if password not in hashes:
            hashes[password] = make_password(password)
        user = User(email=User.objects.normalize_email(email), role='manager' if approver else 'employee',
                    must_change_password=True)
        user.password = hashes[password]
        user.save()
        emp = Employee.objects.create(
            user=user, employee_id=emp_id, p_number=r['p_number'] or None, full_name=r['name'].title() if r['name'].isupper() else r['name'],
            gender='male', joining_date=PLACEHOLDER_JOIN, employment_type='permanent',
            salary_grade=r['cadre'][:50], department=dept, designation=desig, unit_role=r['unit_role'])
        initialize_employee_balances(emp)
        r.update(employee_id=str(emp.pk), login=emp_id, password=password)
        logins.append({'name': emp.full_name, 'unit': dept.name, 'login': emp_id, 'password': password,
                       'place': r['unit_role'] or 'staff'})
    return planned, logins
