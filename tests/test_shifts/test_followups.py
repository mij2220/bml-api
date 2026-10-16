"""Attendance follow-ups and tools (15 Oct 2026)."""
from datetime import date, timedelta
from decimal import Decimal

import pytest

from .conftest import sheet, entry
from .test_attendance import full_flow, D1

API = '/api/v1'


# ---------------------------------------------------------------- 24-day rota
class TestTwentyFourDayRota:

    def set_24(self, plant, client_for):
        admin = client_for(plant.org.admin)
        r = admin.patch(f'{API}/shifts/admin/sections/{plant.section.pk}/',
                        {'pattern': 'M,M,M,M,M,M,O,O,E,E,E,E,E,E,O,O,N,N,N,N,N,N,O,O'}, format='json')
        assert r.status_code == 200, r.data
        r = admin.post(f'{API}/shifts/admin/sections/{plant.section.pk}/spread-groups/', {}, format='json')
        assert r.status_code == 200, r.data
        plant.section.refresh_from_db()
        return r.data['data']['offsets']

    def test_spread_groups_and_every_day_covered(self, plant, client_for):
        assert self.set_24(plant, client_for) == {'A': 0, 'B': 6, 'C': 12, 'D': 18}
        r = client_for(plant.org.admin).get(f'{API}/shifts/roster/', {'section': str(plant.section.pk), 'month': '2026-10'})
        assert r.data['data']['problem_days'] == []

    def test_crew_week_is_each_8_day_block(self, plant, client_for):
        from apps.shifts import roster
        self.set_24(plant, client_for)
        g = plant.section.groups.get(name='A')          # offset 0: anchor day = first Morning
        a = plant.section.pattern_anchor
        assert roster.crew_cycle(plant.section, g, a) == (a, a + timedelta(days=7))
        assert roster.crew_cycle(plant.section, g, a + timedelta(days=7)) == (a, a + timedelta(days=7))
        assert roster.crew_cycle(plant.section, g, a + timedelta(days=9)) == (a + timedelta(days=8), a + timedelta(days=15))
        assert roster.off_day_kind(plant.section, g, a + timedelta(days=6)) == 'FO'
        assert roster.off_day_kind(plant.section, g, a + timedelta(days=7)) == 'SO'

    def test_spread_refused_after_lock_and_for_non_admin(self, plant, client_for):
        assert client_for(plant.manager).post(
            f'{API}/shifts/admin/sections/{plant.section.pk}/spread-groups/', {}, format='json').status_code == 403
        full_flow(plant, D1)
        r = client_for(plant.org.admin).post(f'{API}/shifts/admin/sections/{plant.section.pk}/spread-groups/', {}, format='json')
        assert r.status_code == 400


# ---------------------------------------------------------------- schedule / on duty
class TestSchedule:

    def test_year_schedule_for_attendance_and_leave_supervisors(self, plant, client_for):
        r = client_for(plant.boardman).get(f'{API}/shifts/schedule/', {'section': str(plant.section.pk), 'year': 2026})
        assert r.status_code == 200, r.data
        d = r.data['data']
        assert d['days'] == 365 and len(d['rows']) == 4
        a = next(x for x in d['rows'] if x['group'] == 'A')
        assert len(a['shifts']) == 365
        # Leave Supervisor without an attendance role can also see it
        assert client_for(plant.org.sic).get(f'{API}/shifts/schedule/', {'section': str(plant.section.pk)}).status_code == 200

    def test_operator_outside_section_refused(self, plant, client_for):
        r = client_for(plant.other_marker).get(f'{API}/shifts/schedule/', {'section': str(plant.section.pk)})
        assert r.status_code == 403

    def test_excel_export(self, plant, client_for):
        r = client_for(plant.boardman).get(f'{API}/shifts/schedule/', {'section': str(plant.section.pk), 'year': 2026, 'export': 'xlsx'})
        assert r.status_code == 200 and r['Content-Type'].startswith('application/vnd.openxmlformats')
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames[0] == 'Jan' and len(wb.sheetnames) == 12

    def test_on_duty(self, plant, client_for):
        r = client_for(plant.boardman).get(f'{API}/shifts/on-duty/', {'section': str(plant.section.pk), 'date': str(D1)})
        shifts = {s['code']: s for s in r.data['data']['shifts']}
        assert [g['group'] for g in shifts['M']['groups']] == ['C']
        assert len(shifts['M']['groups'][0]['people']) == 2


# ---------------------------------------------------------------- CD for gazetted holiday worked
@pytest.fixture
def gh_on_d1(plant):
    from apps.leaves.models import HolidayCalendar, PublicHoliday
    cal = HolidayCalendar.objects.create(name='GH 2026', year=2026, is_default=True)
    return PublicHoliday.objects.create(calendar=cal, date=D1, name='Test Holiday')


class TestCdDue:

    def test_listed_after_lock_and_given_once(self, plant, client_for, gh_on_d1):
        sv = client_for(plant.supervisor)
        params = {'section': str(plant.section.pk), 'month': '2026-09'}
        assert sv.get(f'{API}/shifts/cd-due/', params).data['data'] == []      # nothing locked yet
        full_flow(plant, D1)
        rows = sv.get(f'{API}/shifts/cd-due/', params).data['data']
        names = {r['name'] for r in rows}
        assert {'C1 Person', 'D1 Person', 'A1 Person'} <= names and 'B1 Person' not in names  # B is off
        c1 = next(r for r in rows if r['name'] == 'C1 Person')
        r = sv.post(f'{API}/shifts/cd-due/', {'section': str(plant.section.pk), 'employee_id': c1['employee_id'],
                                              'holiday_id': c1['holiday_id']}, format='json')
        assert r.status_code == 201, r.data
        assert 'C1 Person' not in {x['name'] for x in sv.get(f'{API}/shifts/cd-due/', params).data['data']}
        r = sv.post(f'{API}/shifts/cd-due/', {'section': str(plant.section.pk), 'employee_id': c1['employee_id'],
                                              'holiday_id': c1['holiday_id']}, format='json')
        assert r.status_code == 400

    def test_marker_cannot_give(self, plant, client_for, gh_on_d1):
        full_flow(plant, D1)
        r = client_for(plant.boardman).get(f'{API}/shifts/cd-due/', {'section': str(plant.section.pk), 'month': '2026-09'})
        assert r.status_code == 403

    def test_person_not_working_refused(self, plant, client_for, gh_on_d1):
        full_flow(plant, D1)
        r = client_for(plant.supervisor).post(f'{API}/shifts/cd-due/', {
            'section': str(plant.section.pk), 'employee_id': str(plant.people['B1'].pk),
            'holiday_id': str(gh_on_d1.pk)}, format='json')
        assert r.status_code == 400


# ---------------------------------------------------------------- unplanned leave
class TestUnplannedLeave:

    def mark(self, plant, day, emp, code):
        bm = plant.c(plant.boardman)
        d = sheet(plant, bm, day).data['data']
        r = bm.patch(f"{API}/shifts/entries/{entry(d, emp)['id']}/", {'code': code}, format='json')
        assert r.status_code == 200, r.data

    def test_runs_are_grouped_and_raised(self, plant, client_for):
        from apps.leaves.models import LeaveApplication, LeaveBalance
        c1 = plant.people['C1']
        LeaveBalance.objects.create(employee=c1, leave_type=plant.org.CL, year=2026, allocated=10)
        self.mark(plant, D1, c1, 'C')
        self.mark(plant, D1 + timedelta(days=1), c1, 'C')
        sv = client_for(plant.supervisor)
        params = {'section': str(plant.section.pk), 'month': '2026-09'}
        runs = sv.get(f'{API}/shifts/unplanned-leave/', params).data['data']
        run = next(r for r in runs if r['name'] == 'C1 Person')
        assert (run['start'], run['end'], run['days'], run['leave_type']) == ('2026-09-01', '2026-09-02', 2, 'Casual Leave')
        r = sv.post(f'{API}/shifts/unplanned-leave/', {'section': str(plant.section.pk), 'employee_id': str(c1.pk),
                                                       'code': 'C', 'start': run['start'], 'end': run['end']}, format='json')
        assert r.status_code == 201, r.data
        app = LeaveApplication.objects.get(employee=c1)
        assert app.status == 'pending' and app.total_days == 2
        assert all(x['name'] != 'C1 Person' for x in sv.get(f'{API}/shifts/unplanned-leave/', params).data['data'])

    def test_validation_message_passed_on(self, plant, client_for):
        c1 = plant.people['C1']                          # no casual leave balance
        self.mark(plant, D1, c1, 'C')
        r = client_for(plant.supervisor).post(f'{API}/shifts/unplanned-leave/', {
            'section': str(plant.section.pk), 'employee_id': str(c1.pk), 'code': 'C',
            'start': str(D1), 'end': str(D1)}, format='json')
        assert r.status_code == 400 and 'could not be raised' in r.data['message']

    def test_marker_cannot_raise(self, plant, client_for):
        r = client_for(plant.boardman).get(f'{API}/shifts/unplanned-leave/', {'section': str(plant.section.pk)})
        assert r.status_code == 403


# ---------------------------------------------------------------- calculator
class TestCalculator:

    def test_all_sop_examples_match(self, plant, client_for):
        r = client_for(plant.accounts).get(f'{API}/overtime/calculator/')
        assert r.status_code == 200
        bad = [e['title'] for e in r.data['data']['examples'] if not e['matches_sop']]
        assert bad == []

    def test_case(self, plant, client_for):
        r = client_for(plant.manager).post(f'{API}/overtime/calculator/', {
            'day_type': 'SO', 'kind': 'normal', 'hours': 10, 'shift': 'N'}, format='json')
        d = r.data['data']
        assert (d['single'], d['double'], d['equivalent_single']) == ('8', '2', '12')
        assert d['night_premium_hours'] == '10' and 'Second Day Off' in d['explanation']

    def test_bad_input_and_no_access(self, plant, client_for):
        c = client_for(plant.manager)
        assert c.post(f'{API}/overtime/calculator/', {'day_type': 'X', 'hours': 2}, format='json').status_code == 400
        assert c.post(f'{API}/overtime/calculator/', {'day_type': 'FO', 'hours': 'NaN'}, format='json').status_code == 400
        assert client_for(plant.org.op_b).get(f'{API}/overtime/calculator/').status_code == 403


# ---------------------------------------------------------------- manpower, notifications, alerts, to-do
class TestManpowerAndNotices:

    def test_shortfall_on_sheet(self, plant, client_for):
        t = plant.section.shift_types.get(code='M')
        t.min_staff = 3
        t.save()
        d = sheet(plant, client_for(plant.boardman), D1).data['data']
        m = next(s for s in d['shifts'] if s['code'] == 'M')
        assert (m['counts']['min_staff'], m['counts']['short']) == (3, 1)

    def test_send_back_notifies_marker(self, plant, client_for):
        from apps.notifications.models import Notification
        bm = client_for(plant.boardman)
        d = sheet(plant, bm, D1).data['data']
        bm.post(f"{API}/shifts/sheets/{d['id']}/submit/", {'shift': 'M'}, format='json')
        r = client_for(plant.supervisor).post(f"{API}/shifts/sheets/{d['id']}/send-back/",
                                              {'shift': 'M', 'remarks': 'Check C2'}, format='json')
        assert r.status_code == 200, r.data
        assert Notification.objects.filter(recipient=plant.boardman, title__startswith='Shift sent back').exists()

    def test_correction_notifies_manager_and_requester(self, plant, client_for):
        from apps.notifications.models import Notification
        sid = full_flow(plant, D1)
        r = client_for(plant.boardman).post(f'{API}/shifts/sheets/{sid}/corrections/', {
            'kind': 'code', 'employee': str(plant.people['C1'].pk), 'code': 'A', 'reason': 'Was absent'}, format='json')
        assert r.status_code == 201, r.data
        assert Notification.objects.filter(recipient=plant.manager, title='Attendance correction waiting').exists()
        client_for(plant.manager).post(f"{API}/shifts/corrections/{r.data['data']['id']}/decide/", {'approve': True}, format='json')
        assert Notification.objects.filter(recipient=plant.boardman, title='Correction approved').exists()

    def test_long_day_alert_in_statement(self, plant, client_for):
        bm = client_for(plant.boardman)
        d = sheet(plant, bm, D1).data['data']
        r = bm.post(f"{API}/shifts/sheets/{d['id']}/extras/", {
            'employee': str(plant.people['C1'].pk), 'hours': 9, 'shift_code': 'E', 'kind': 'normal'}, format='json')
        assert r.status_code in (200, 201), r.data
        full_flow(plant, D1)
        rows = client_for(plant.accounts).get(f'{API}/overtime/statement/', {
            'section': str(plant.section.pk), 'month': '2026-09'}).data['data']['rows']
        c1 = next(x for x in rows if x['name'] == 'C1 Person')
        assert any('over 16' in a for a in c1['alerts']), c1['alerts']

    def test_todo_counts(self, plant, client_for):
        from django.utils import timezone
        full_flow(plant, timezone.localdate() - timedelta(days=2), lock=False)
        t = client_for(plant.manager).get(f'{API}/shifts/todo/').data['data']
        assert t['to_lock'] == 1
        assert client_for(plant.boardman).get(f'{API}/shifts/todo/').data['data']['to_lock'] == 0


# ---------------------------------------------------------------- Engro 2026 schedule
# Engro's "Pattern" sheet (Shift_Schedule_2026.xlsx): cycle day -> crew on Morning, Evening, Night, Off
ENGRO_CYCLE = ['ABCD', 'ABDC', 'ABDC', 'ACDB', 'ACDB', 'BCDA', 'BCDA', 'BCAD', 'BCAD', 'BDAC', 'BDAC', 'CDAB',
               'CDAB', 'CDBA', 'CDBA', 'CABD', 'CABD', 'DABC', 'DABC', 'DACB', 'DACB', 'DBCA', 'DBCA', 'ABCD']


class TestEngro2026Schedule:

    def test_every_day_of_2026_matches_engro_pdf(self, plant, client_for):
        from apps.shifts import roster
        r = client_for(plant.org.admin).post(f'{API}/shifts/admin/sections/{plant.section.pk}/spread-groups/',
                                             {'preset': 'engro_2026'}, format='json')
        assert r.status_code == 200, r.data
        plant.section.refresh_from_db()
        groups = {g.name: g for g in plant.section.groups.all()}
        d = date(2026, 1, 1)
        while d.year == 2026:
            want = ENGRO_CYCLE[(d - date(2026, 1, 1)).days % 24]
            for code, crew in zip('MENO', want):
                assert roster.shift_on(plant.section, groups[crew], d) == code, (d, crew, code)
            d += timedelta(days=1)
        # spot checks from the Daily Schedule sheet
        assert roster.shift_on(plant.section, groups['A'], date(2026, 12, 31)) == 'M'
        assert roster.shift_on(plant.section, groups['C'], date(2026, 12, 31)) == 'E'
